

# ===== ATTENDANCE OUTPUT GENERATOR =====
#
# Generates an attendance-marked result file from the session's
# student list. The output preserves every student from the original
# uploaded list (Roll + Name) and adds an "Attendance" column
# (Present / Absent).
#
# Supported output formats:
#   - .xlsx  (openpyxl)
#   - .docx  (python-docx, a Word table)
#   - .csv   (stdlib csv)
#   - .xls   (legacy Excel is read but not written; generated as .xlsx)

import csv
import io
import os


def _rows_for(students_with_status):
    """Return ordered [ [roll, name, status], ... ] rows.

    students_with_status: list of dicts with {roll, name, status}.
    Ordered by roll for a stable, predictable result file.
    """
    ordered = sorted(
        students_with_status,
        key=lambda s: int(s.get("roll") or 0)
    )

    rows = []

    for s in ordered:

        rows.append([
            s.get("roll"),
            s.get("name") or "",
            (s.get("status") or "Absent").capitalize()
        ])

    return rows


def _suggest_name(base_name, fmt):
    """Produce a result filename: '<base>_attendance.<fmt>'."""

    root = os.path.splitext(base_name or "attendance")[0]
    root = root.strip().replace(" ", "_")

    if not root:
        root = "attendance"

    return root + "_attendance." + fmt


def generate_attendance_file(students_with_status, fmt, base_name=""):
    """Generate an attendance-marked file.

    Returns (bytes_io, filename, mimetype).
    fmt is one of: xlsx, docx, xls, csv.
    """

    fmt = (fmt or "xlsx").lower().lstrip(".")

    if fmt == "docx":
        return _generate_docx(students_with_status, base_name)

    if fmt == "csv":
        return _generate_csv(students_with_status, base_name)

    # xlsx, xls, or anything else -> Excel XLSX output.
    return _generate_xlsx(students_with_status, base_name)


def _generate_xlsx(students_with_status, base_name=""):
    """Generate an .xlsx result file with Roll | Name | Attendance."""

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    rows = _rows_for(students_with_status)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    header_font = Font(bold=True)
    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    headers = ["Roll No", "Name", "Attendance"]

    sheet.append(headers)

    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        sheet.append(row)

    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 14

    sheet.freeze_panes = "A2"

    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)

    filename = _suggest_name(base_name, "xlsx")

    return buf, filename, (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    )


def _generate_docx(students_with_status, base_name=""):
    """Generate a .docx result file with a Word table."""

    from docx import Document
    from docx.shared import RGBColor

    rows = _rows_for(students_with_status)

    document = Document()

    document.add_heading("Attendance Record", level=1)

    table = document.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"

    hdr = table.rows[0].cells
    hdr[0].text = "Roll No"
    hdr[1].text = "Name"
    hdr[2].text = "Attendance"

    for roll, name, status in rows:

        cells = table.add_row().cells
        cells[0].text = str(roll)
        cells[1].text = str(name)
        cells[2].text = status

    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    filename = _suggest_name(base_name, "docx")

    return buf, filename, (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    )


def _generate_csv(students_with_status, base_name=""):
    """Generate a .csv result file with Roll No, Name, Attendance."""

    rows = _rows_for(students_with_status)

    buf = io.StringIO()
    writer = csv.writer(buf)

    writer.writerow(["Roll No", "Name", "Attendance"])

    for row in rows:
        writer.writerow(row)

    data = buf.getvalue().encode("utf-8-sig")

    filename = _suggest_name(base_name, "csv")

    return io.BytesIO(data), filename, "text/csv"
