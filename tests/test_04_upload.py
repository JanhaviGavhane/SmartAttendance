"""TEST 4 - Student file upload: CSV/XLSX/DOCX, invalid, empty, dupes, 1-100."""

import io
import csv
import os

from tests.conftest import register


def _auth(client):
    register(client, "uploader")


def _csv_bytes(headers, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _upload(client, name, content_bytes, mode="replace"):
    return client.post(
        "/students/upload",
        data={
            "file": (io.BytesIO(content_bytes), name),
            "mode": mode,
        },
        content_type="multipart/form-data",
    )


def test_csv_upload(client):
    _auth(client)
    content = _csv_bytes(
        ["Roll No", "Name"],
        [[1, "Student One"], [2, "Student Two"]],
    )
    resp = _upload(client, "students.csv", content)
    data = resp.get_json()
    assert resp.status_code == 200
    assert data["success"] is True
    assert len(data["students"]) == 2


def test_xlsx_upload(client):
    _auth(client)
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Roll", "Name"])
    ws.append([1, "X One"])
    ws.append([2, "X Two"])
    buf = io.BytesIO()
    wb.save(buf)
    resp = _upload(client, "students.xlsx", buf.getvalue())
    data = resp.get_json()
    assert data["success"] is True
    assert len(data["students"]) == 2


def test_docx_upload(client):
    _auth(client)
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Roll"
    table.cell(0, 1).text = "Name"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "D One"
    buf = io.BytesIO()
    doc.save(buf)
    resp = _upload(client, "students.docx", buf.getvalue())
    data = resp.get_json()
    assert data["success"] is True
    assert len(data["students"]) == 1


def test_invalid_file_type_rejected(client):
    _auth(client)
    resp = _upload(client, "students.txt", b"not at all a student file")
    assert resp.status_code == 400
    assert resp.get_json()["success"] is False


def test_empty_file_rejected(client):
    _auth(client)
    resp = _upload(client, "students.csv", b"")
    assert resp.status_code == 400


def test_duplicate_rolls_skipped(client):
    _auth(client)
    content = _csv_bytes(
        ["Roll", "Name"],
        [[1, "First"], [1, "Duplicate"]],
    )
    resp = _upload(client, "dup.csv", content)
    data = resp.get_json()
    assert data["success"] is True
    assert len(data["students"]) == 1
    assert len(data["duplicates"]) == 1


def test_rolls_1_to_100(client):
    _auth(client)
    rows = [[i, "Student %d" % i] for i in range(1, 101)]
    content = _csv_bytes(["Roll", "Name"], rows)
    resp = _upload(client, "big.csv", content)
    data = resp.get_json()
    assert data["success"] is True
    assert len(data["students"]) == 100
    rolls = [s["roll"] for s in data["students"]]
    assert rolls == list(range(1, 101))